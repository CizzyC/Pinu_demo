from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

path = "/Users/cizzyc/Documents/Pinu/Pinu_demo/outputs/level_xp_model/Pinu等级经验曲线_简版.xlsx"

wb = load_workbook(path)
ws = wb["日期-等级经验曲线"]

for drawing in list(ws._charts):
    ws._charts.remove(drawing)

chart = LineChart()
chart.title = "留存周期 - 等级/累计经验曲线"
chart.style = 13
chart.y_axis.title = "等级"
chart.x_axis.title = "相对天数"
chart.height = 12
chart.width = 24

days = Reference(ws, min_col=2, min_row=4, max_row=369)
level_data = Reference(ws, min_col=5, min_row=3, max_row=369)
chart.add_data(level_data, titles_from_data=True)
chart.set_categories(days)

xp_chart = LineChart()
xp_data = Reference(ws, min_col=4, min_row=3, max_row=369)
xp_chart.add_data(xp_data, titles_from_data=True)
xp_chart.set_categories(days)
xp_chart.y_axis.axId = 200
xp_chart.y_axis.title = "累计经验"
xp_chart.y_axis.crosses = "max"

chart += xp_chart
chart.legend.position = "b"
ws.add_chart(chart, "G15")

daily_chart = LineChart()
daily_chart.title = "当日 XP 随留存周期增长"
daily_chart.style = 12
daily_chart.y_axis.title = "当日 XP"
daily_chart.x_axis.title = "相对天数"
daily_chart.height = 8
daily_chart.width = 24
daily_xp_data = Reference(ws, min_col=3, min_row=3, max_row=369)
daily_chart.add_data(daily_xp_data, titles_from_data=True)
daily_chart.set_categories(days)
daily_chart.legend = None
ws.add_chart(daily_chart, "G39")

wb.save(path)
